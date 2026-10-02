"""Integration tests: statistics from the events table and the Excel/PDF exports."""

import io
from datetime import date

import pytest

from app.db.sqlite import connection
from app.services import reports


def add_event(type_, created_at, visitor="v1", query=None, result=None, latency=None):
    with connection() as conn:
        conn.execute(
            """INSERT INTO events (type, visitor_hash, query, result, latency_ms, created_at)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (type_, visitor, query, result, latency, created_at),
        )


@pytest.fixture
def activity(client):
    # Times are UTC; reports group by Saudi time (UTC+3).
    add_event("visit", "2026-09-01 10:00:00", visitor="a")
    add_event("visit", "2026-09-01 11:00:00", visitor="a")
    add_event("visit", "2026-09-01 22:30:00", visitor="b")  # 01:30 on Sep 2 in Saudi time
    add_event("search", "2026-09-01 10:01:00", visitor="a", query="إنما الأعمال", result="verified", latency=100)
    add_event("search", "2026-09-01 10:05:00", visitor="a", query="إنما الأعمال", result="verified", latency=300)
    add_event("search", "2026-09-01 11:00:00", visitor="a", query="اطلبوا العلم", result="no_match", latency=200)
    add_event("chat", "2026-09-01 10:02:00", visitor="a", query="ما حكمه؟")
    add_event("visit", "2026-08-15 10:00:00", visitor="c")  # outside the range below


def test_summary_counts_and_groups_by_local_day(activity):
    s = reports.build_summary(date(2026, 9, 1), date(2026, 9, 2))

    assert s.visits == 3
    assert s.unique_visitors == 2
    assert s.searches == 3
    assert s.chats == 1
    assert s.avg_search_latency_ms == 200
    assert [(d.day, d.visits) for d in s.daily] == [("2026-09-01", 2), ("2026-09-02", 1)]
    assert s.top_searches[0].label == "إنما الأعمال" and s.top_searches[0].count == 2
    assert s.top_questions[0].label == "ما حكمه؟"
    assert {i.label: i.count for i in s.search_results} == {"موثّق": 2, "لا تطابق": 1}


def test_empty_days_are_filled_with_zeros(client):
    s = reports.build_summary(date(2026, 9, 1), date(2026, 9, 3))
    assert [d.visits for d in s.daily] == [0, 0, 0]


def test_excel_export(admin_client, activity):
    from openpyxl import load_workbook

    admin_client.post("/api/sources", files={"file": ("كتاب.txt", "نص عربي طويل بما يكفي ليكون مقطعًا مفهرسًا في القاعدة".encode())})
    res = admin_client.get("/api/reports/export?format=xlsx&from=2026-09-01&to=2026-09-02")

    assert res.status_code == 200
    assert "isnad-report-2026-09-01-to-2026-09-02.xlsx" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["الملخص", "يومي", "أكثر عمليات البحث", "أكثر أسئلة الحوار", "نتائج البحث", "الملفات"]
    summary = {row[0]: row[1] for row in wb["الملخص"].iter_rows(min_row=2, values_only=True)}
    assert summary["عدد الزيارات"] == 3
    assert wb["الملفات"]["A2"].value == "كتاب.txt"
    assert wb["الملخص"].sheet_view.rightToLeft


def test_pdf_export(admin_client, activity):
    res = admin_client.get("/api/reports/export?format=pdf&from=2026-09-01&to=2026-09-02")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")


def test_pdf_export_with_no_data(admin_client):
    res = admin_client.get("/api/reports/export?format=pdf&from=2026-01-01&to=2026-01-02")
    assert res.status_code == 200
    assert res.content.startswith(b"%PDF")


@pytest.fixture
def hostile_activity(client):
    """What a visitor can type into the search box ends up in the admin's reports."""
    add_event("search", "2026-09-01 10:00:00", query='=HYPERLINK("http://example.com","افتح")')
    add_event("chat", "2026-09-01 10:01:00", query="<img src='http://127.0.0.1:9/x.png'/>")
    add_event("chat", "2026-09-01 10:02:00", query="<b>نص لم يُغلق")
    add_event("chat", "2026-09-01 10:03:00", query="أ & ب < ج")


def test_excel_keeps_visitor_text_as_text(admin_client, hostile_activity):
    from openpyxl import load_workbook

    res = admin_client.get("/api/reports/export?format=xlsx&from=2026-09-01&to=2026-09-01")
    cell = load_workbook(io.BytesIO(res.content))["أكثر عمليات البحث"]["A2"]
    assert cell.value == '=HYPERLINK("http://example.com","افتح")'
    assert cell.data_type == "s"  # not a formula


def test_pdf_shows_visitor_markup_as_text(admin_client, hostile_activity, monkeypatch):
    def no_fetch(*args, **kwargs):
        raise AssertionError("the PDF report tried to load an image named in a visitor's text")

    monkeypatch.setattr("reportlab.lib.utils.ImageReader.__init__", no_fetch)
    res = admin_client.get("/api/reports/export?format=pdf&from=2026-09-01&to=2026-09-01")
    assert res.status_code == 200
    assert res.content.startswith(b"%PDF")
