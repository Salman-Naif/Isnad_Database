"""Tests for text extraction from each supported file type."""

import io
import shutil

import pytest

from app.config import get_settings
from app.services.extraction import ExtractionError, extract_text
from app.services.ocr import PageResult

ARABIC = "إنما الأعمال بالنيات"


def test_txt_utf8():
    assert extract_text(ARABIC.encode("utf-8"), "a.txt").text == ARABIC


def test_txt_utf8_bom():
    assert extract_text(ARABIC.encode("utf-8-sig"), "a.txt").text == ARABIC


def test_txt_windows_arabic_encoding():
    assert extract_text(ARABIC.encode("cp1256"), "a.txt").text == ARABIC


def test_docx_paragraphs_and_tables():
    from docx import Document

    doc = Document()
    doc.add_paragraph(ARABIC)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "الحكم"
    table.cell(0, 1).text = "صحيح"
    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text(buf.getvalue(), "a.docx").text
    assert ARABIC in text
    assert "الحكم | صحيح" in text


def test_unsupported_type():
    with pytest.raises(ExtractionError):
        extract_text(b"data", "a.exe")


def test_corrupt_pdf():
    with pytest.raises(ExtractionError):
        extract_text(b"not a pdf", "a.pdf")


def _tesseract_available() -> bool:
    return bool(get_settings().tesseract_cmd or shutil.which("tesseract"))


@pytest.mark.skipif(not _tesseract_available(), reason="Tesseract is not installed")
def test_scanned_image_ocr():
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (400, 80), "white")
    ImageDraw.Draw(image).text((10, 25), "ISNAD OCR TEST", fill="black")
    buf = io.BytesIO()
    image.save(buf, format="PNG")

    result = extract_text(buf.getvalue(), "scan.png")
    assert result.ocr_pages == 1
    assert "ISNAD" in result.text.upper()


# --- Garbled text layers and OCR quality ---

from app.services import extraction  # noqa: E402

# Real junk from a scanned hadith book: Arabic interleaved with misread Latin letters.
GARBLED = "دعا BE gall أن التعال بعد رسول الله وي أفضل من جعفر INKS sect enoncacnupicce غريب أنا دار الحكمة AGN sit LE Snes ee ee"
CLEAN_ARABIC = "إنما الأعمال بالنيات وإنما لكل امرئ ما نوى فمن كانت هجرته إلى الله ورسوله"


def test_garbled_ocr_is_detected():
    assert extraction.looks_garbled(GARBLED)
    assert not extraction.looks_garbled(CLEAN_ARABIC)
    assert not extraction.looks_garbled("An English page of text is not garbled at all, only Latin.")
    assert not extraction.looks_garbled("قصير BE")  # too little text to judge


def test_presentation_forms_become_normal_letters():
    # "ﻻ ﺇﻟﻪ ﺇﻻ ﺍﻟﻠﻪ" as PDFs often store it (presentation forms)
    assert extraction.normalize_forms("ﻻ ﺇﻟﻪ ﺇﻻ ﺍﻟﻠﻪ") == "لا إله إلا الله"


def test_broken_text_layers_are_detected():
    spaced = " ".join("انما الاعمال بالنيات وانما لكل امرئ ما نوى فمن كانت هجرته الى الله ورسوله")
    backwards = " ".join(w[::-1] for w in (CLEAN_ARABIC + " والحديث الصحيح الاسناد المتصل والرواة العدول الضابطين").split())
    assert extraction.looks_broken(spaced)
    assert extraction.looks_broken(backwards)
    assert not extraction.looks_broken(CLEAN_ARABIC + " " + CLEAN_ARABIC)
    assert extraction.text_layer_usable(CLEAN_ARABIC)
    assert not extraction.text_layer_usable(GARBLED)
    assert not extraction.text_layer_usable("١٢")  # page number only: a scan


# --- PDFs: text layer kept, scanned pages sent to OCR ---


class FakeReader:
    """Stands in for the OCR engine: returns a fixed reading for every image."""

    reads = 0

    def __init__(self, text=CLEAN_ARABIC, **flags):
        self.text, self.flags = text, flags

    def __call__(self):
        return self

    def __enter__(self):
        return self

    def __exit__(self, *_):
        pass

    def read(self, image):
        FakeReader.reads += 1
        return PageResult(self.text, **self.flags)


@pytest.fixture
def fake_ocr(monkeypatch):
    def install(text=CLEAN_ARABIC, **flags):
        reader = FakeReader(text, **flags)
        monkeypatch.setattr(extraction, "PageReader", reader)
        FakeReader.reads = 0
        return reader

    return install


def _text_pdf(text: str) -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    pdf = canvas.Canvas(buf)
    pdf.drawString(72, 720, text)
    pdf.showPage()
    pdf.save()
    return buf.getvalue()


def _scanned_pdf(pages: int = 1) -> bytes:
    from PIL import Image, ImageDraw

    images = []
    for _ in range(pages):
        image = Image.new("L", (850, 1100), "white")
        draw = ImageDraw.Draw(image)
        for y in range(150, 1000, 40):
            draw.rectangle((100, y, 750, y + 18), fill="black")  # lines of "text"
        images.append(image)
    buf = io.BytesIO()
    images[0].save(buf, format="PDF", save_all=True, append_images=images[1:], resolution=100)
    return buf.getvalue()


def test_pdf_text_layer_is_used_without_ocr(fake_ocr):
    fake_ocr("should not be used")
    result = extract_text(_text_pdf("A real text layer with plenty of words in it"), "book.pdf")
    assert "A real text layer" in result.text
    assert result.ocr_pages == 0 and FakeReader.reads == 0


def test_scanned_pdf_pages_are_ocrd_in_order(fake_ocr):
    fake_ocr()
    progress = []
    result = extract_text(_scanned_pdf(3), "scan.pdf", progress.append)
    assert result.ocr_pages == 3
    assert result.text == "\n\n".join([CLEAN_ARABIC] * 3)
    assert progress[-1] == "قراءة الصفحات 3/3 — منها 3 بالتعرّف الضوئي"


def test_pages_worth_a_look_are_noted(fake_ocr):
    fake_ocr(disputed=True)
    result = extract_text(_scanned_pdf(2), "scan.pdf")
    assert result.note and "1، 2" in result.note


def test_ocr_finding_nothing_keeps_the_layer():
    assert extraction._choose("", GARBLED) == GARBLED
    assert extraction._choose(CLEAN_ARABIC, GARBLED) == CLEAN_ARABIC


def test_worse_tesseract_reading_does_not_replace_a_garbled_layer():
    # Without a vision model (no key in tests), a worse Tesseract reading loses.
    worse = "BE gall INKS sect enoncacnupicce AGN sit LE Snes ee ee و"
    assert extraction._choose(worse, GARBLED) == GARBLED


def test_images_and_multipage_tiffs(fake_ocr):
    from PIL import Image

    fake_ocr()
    frames = [Image.new("L", (300, 300), "white") for _ in range(2)]
    buf = io.BytesIO()
    frames[0].save(buf, format="TIFF", save_all=True, append_images=frames[1:])
    result = extract_text(buf.getvalue(), "scan.tif")
    assert result.ocr_pages == 2


def test_scans_pasted_into_word_are_ocrd(fake_ocr):
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    fake_ocr("نص الصفحة الممسوحة")
    page, icon = io.BytesIO(), io.BytesIO()
    Image.new("L", (1200, 1600), "white").save(page, format="PNG")
    Image.new("L", (64, 64), "white").save(icon, format="PNG")
    doc = Document()
    doc.add_paragraph("مقدمة الكتاب")
    doc.add_picture(page, width=Inches(3))
    doc.add_picture(icon, width=Inches(0.3))  # a logo: not OCR'd
    doc.add_paragraph("خاتمة")
    buf = io.BytesIO()
    doc.save(buf)

    result = extract_text(buf.getvalue(), "book.docx")
    assert result.text == "مقدمة الكتاب\n\nنص الصفحة الممسوحة\n\nخاتمة"
    assert result.ocr_pages == 1


def test_ocr_uses_arabic_only_by_default():
    assert get_settings().ocr_languages == "ara"
