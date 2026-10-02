"""
Statistics and reports: visits, visitors, searches and chat questions reported by the
main website, plus the uploaded sources — for a date range, exportable to Excel and PDF.

Events are stored in UTC and grouped by day in REPORT_UTC_OFFSET_HOURS (Saudi time by default).
"""

import glob
import io
import logging
import re
from datetime import UTC, date, datetime, timedelta
from functools import lru_cache
from html import escape

from app.config import get_settings
from app.db.sqlite import connection
from app.models.schemas import CountItem, DailyStat, ReportSummary, SourceInfo

MAX_RANGE_DAYS = 366
TOP_N = 10

RESULT_LABELS = {
    "verified": "موثّق",
    "found": "موجود بلا حكم",
    "distorted": "تحريف محتمل",
    "no_match": "لا تطابق",
}
KIND_LABELS = {"document": "مستند", "structured_hadith": "أحاديث منظّمة"}

logger = logging.getLogger(__name__)


class ReportError(ValueError):
    """Invalid report request (message is shown to the admin)."""


def default_range() -> tuple[date, date]:
    """The last 30 days, ending today in report time."""
    today = (datetime.now(UTC) + _offset()).date()
    return today - timedelta(days=29), today


def _offset() -> timedelta:
    return timedelta(hours=get_settings().report_utc_offset_hours)


# Fixed SQL statements: every value (the time offset and the dates) is bound as a ? parameter.
# created_at / uploaded_at are UTC 'YYYY-MM-DD HH:MM:SS'; datetime(x, '+3 hours') shifts them
# to report time before date() takes the day.
_TOTALS_SQL = """
    SELECT SUM(type = 'visit')  AS visits,
           SUM(type = 'search') AS searches,
           SUM(type = 'chat')   AS chats,
           COUNT(DISTINCT visitor_hash) AS visitors,
           AVG(CASE WHEN type = 'search' THEN latency_ms END) AS avg_latency
    FROM events WHERE date(datetime(created_at, ?)) BETWEEN ? AND ?"""
_BY_DAY_SQL = """
    SELECT date(datetime(created_at, ?)) AS day,
           SUM(type = 'visit')  AS visits,
           SUM(type = 'search') AS searches,
           SUM(type = 'chat')   AS chats,
           COUNT(DISTINCT visitor_hash) AS visitors
    FROM events WHERE date(datetime(created_at, ?)) BETWEEN ? AND ? GROUP BY day"""
_TOP_QUERIES_SQL = """
    SELECT query, COUNT(*) AS n FROM events
    WHERE type = ? AND query IS NOT NULL AND date(datetime(created_at, ?)) BETWEEN ? AND ?
    GROUP BY query ORDER BY n DESC, query LIMIT ?"""
_RESULTS_SQL = """
    SELECT result, COUNT(*) AS n FROM events
    WHERE type = 'search' AND result IS NOT NULL AND date(datetime(created_at, ?)) BETWEEN ? AND ?
    GROUP BY result ORDER BY n DESC"""
_UPLOADS_SQL = """
    SELECT COUNT(*) FROM sources
    WHERE status = 'ready' AND date(datetime(uploaded_at, ?)) BETWEEN ? AND ?"""


def _offset_modifier() -> str:
    """SQLite datetime() modifier for the report time zone, e.g. '+3 hours'."""
    return f"{get_settings().report_utc_offset_hours:+d} hours"


def build_summary(date_from: date, date_to: date) -> ReportSummary:
    if date_from > date_to:
        raise ReportError("تاريخ البداية بعد تاريخ النهاية")
    if (date_to - date_from).days >= MAX_RANGE_DAYS:
        raise ReportError(f"أقصى مدة للتقرير {MAX_RANGE_DAYS} يومًا")

    shift = _offset_modifier()
    in_period = (shift, date_from.isoformat(), date_to.isoformat())

    with connection() as conn:
        totals = conn.execute(_TOTALS_SQL, in_period).fetchone()
        by_day = {r["day"]: r for r in conn.execute(_BY_DAY_SQL, (shift, *in_period))}

        def top(event_type: str) -> list[CountItem]:
            rows = conn.execute(_TOP_QUERIES_SQL, (event_type, *in_period, TOP_N))
            return [CountItem(label=r["query"], count=r["n"]) for r in rows]

        results = [
            CountItem(label=RESULT_LABELS.get(r["result"], r["result"]), count=r["n"])
            for r in conn.execute(_RESULTS_SQL, in_period)
        ]

        files, chunks = conn.execute(
            "SELECT COUNT(*), COALESCE(SUM(chunks), 0) FROM sources WHERE status = 'ready'"
        ).fetchone()
        uploads = conn.execute(_UPLOADS_SQL, in_period).fetchone()[0]

        top_searches, top_questions = top("search"), top("chat")

    daily = []
    current = date_from
    while current <= date_to:
        row = by_day.get(current.isoformat())
        daily.append(
            DailyStat(
                day=current.isoformat(),
                visits=row["visits"] if row else 0,
                visitors=row["visitors"] if row else 0,
                searches=row["searches"] if row else 0,
                chats=row["chats"] if row else 0,
            )
        )
        current += timedelta(days=1)

    return ReportSummary(
        date_from=date_from.isoformat(),
        date_to=date_to.isoformat(),
        visits=totals["visits"] or 0,
        unique_visitors=totals["visitors"] or 0,
        searches=totals["searches"] or 0,
        chats=totals["chats"] or 0,
        avg_search_latency_ms=round(totals["avg_latency"]) if totals["avg_latency"] else None,
        daily=daily,
        top_searches=top_searches,
        top_questions=top_questions,
        search_results=results,
        files=files,
        chunks=chunks,
        uploads_in_period=uploads,
    )


def _summary_rows(s: ReportSummary) -> list[tuple[str, object]]:
    return [
        ("الفترة", f"من {s.date_from} إلى {s.date_to}"),
        ("عدد الزيارات", s.visits),
        ("عدد الزوار (فريد)", s.unique_visitors),
        ("عمليات البحث", s.searches),
        ("أسئلة الحوار", s.chats),
        ("متوسط زمن البحث (ms)", s.avg_search_latency_ms if s.avg_search_latency_ms is not None else "—"),
        ("الملفات في قاعدة البيانات", s.files),
        ("المقاطع المفهرسة", s.chunks),
        ("ملفات رُفعت خلال الفترة", s.uploads_in_period),
    ]


def _size(num_bytes: int) -> str:
    if num_bytes >= 1024 * 1024:
        return f"{num_bytes / 1024 / 1024:.1f} MB"
    return f"{max(1, round(num_bytes / 1024))} KB"


# --- Excel ---


def to_excel(summary: ReportSummary, sources: list[SourceInfo]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0B6E4F")

    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title: str, headers: list[str], rows: list[list[object]], widths: list[int]):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                # Visitors' searches and file names are text: one starting with "=" must
                # not become a formula in the admin's Excel (=HYPERLINK, =WEBSERVICE…).
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
        for i, width in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + i)].width = width
        ws.freeze_panes = "A2"

    sheet("الملخص", ["البند", "القيمة"], [list(r) for r in _summary_rows(summary)], [30, 28])
    sheet(
        "يومي",
        ["اليوم", "الزيارات", "الزوار", "عمليات البحث", "أسئلة الحوار"],
        [[d.day, d.visits, d.visitors, d.searches, d.chats] for d in summary.daily],
        [14, 12, 12, 14, 14],
    )
    sheet("أكثر عمليات البحث", ["النص", "العدد"], [[i.label, i.count] for i in summary.top_searches], [70, 10])
    sheet("أكثر أسئلة الحوار", ["السؤال", "العدد"], [[i.label, i.count] for i in summary.top_questions], [70, 10])
    sheet("نتائج البحث", ["النتيجة", "العدد"], [[i.label, i.count] for i in summary.search_results], [24, 10])
    sheet(
        "الملفات",
        ["اسم الملف", "النوع", "الحجم", "المقاطع", "صفحات OCR", "رفعه", "تاريخ الرفع (UTC)"],
        [
            [s.filename, KIND_LABELS.get(s.kind, s.kind), _size(s.size_bytes), s.chunks,
             s.ocr_pages, s.uploaded_by or "", s.uploaded_at]
            for s in sources
        ],
        [40, 16, 12, 10, 12, 14, 20],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- PDF ---

_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",
    "C:/Windows/Fonts/tahoma.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
]


@lru_cache
def _pdf_font() -> str:
    """Register a TTF font that has Arabic glyphs and return its name."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    configured = get_settings().report_font_path
    candidates = ([configured] if configured else []) + _FONT_CANDIDATES
    candidates += sorted(glob.glob("/usr/share/fonts/**/*Arabic*-Regular.ttf", recursive=True))
    for path in candidates:
        try:
            pdfmetrics.registerFont(TTFont("ReportArabic", path))
            return "ReportArabic"
        except Exception as exc:  # missing or unreadable font file: try the next one
            logger.debug("Report font %s not usable: %s", path, exc)
    raise ReportError("لا يوجد خط عربي لإنشاء PDF — اضبط REPORT_FONT_PATH على ملف خط TTF يدعم العربية")


# Dates and times inside Arabic text; the bidi algorithm would otherwise flip
# "2026-09-01" to "01-09-2026", so each one is wrapped in a left-to-right embedding.
_DATE_PATTERN = re.compile(r"\d{4}-\d{2}-\d{2}(?: \d{2}:\d{2}(?::\d{2})?)?")


def _ar(text: object) -> str:
    """Shape Arabic letters and reorder right-to-left, which reportlab can't do itself.

    Escaped for reportlab's Paragraph, which reads its text as markup: the report holds what
    visitors typed, and \u00ab<img src=\u2026>\u00bb there would make the server fetch it, \u00ab<b>\u00bb break the PDF.
    """
    import arabic_reshaper
    from bidi.algorithm import get_display

    text = _DATE_PATTERN.sub(lambda m: f"\u202a{m.group()}\u202c", str(text))
    return escape(get_display(arabic_reshaper.reshape(text)), quote=False)


def to_pdf(summary: ReportSummary, sources: list[SourceInfo]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _pdf_font()
    title = ParagraphStyle("t", fontName=font, fontSize=18, alignment=TA_RIGHT, leading=26,
                           textColor=colors.HexColor("#07291F"))
    heading = ParagraphStyle("h", fontName=font, fontSize=13, alignment=TA_RIGHT, leading=20,
                             textColor=colors.HexColor("#0B6E4F"), spaceBefore=12, spaceAfter=6)
    cell = ParagraphStyle("c", fontName=font, fontSize=9, alignment=TA_RIGHT, leading=13)
    note = ParagraphStyle("n", parent=cell, textColor=colors.HexColor("#5C6B67"))

    def table(headers: list[str], rows: list[list[object]], widths: list[float]) -> Table:
        # Columns are reversed so the first column appears on the right (RTL reading order).
        data = [[Paragraph(_ar(v), cell) for v in reversed(row)] for row in [headers, *rows]]
        t = Table(data, colWidths=list(reversed(widths)), repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3F0EA")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D8D0")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    def short(value: object, limit: int = 80) -> object:
        # reportlab wraps bidi text in the wrong line order, so keep cells to one line.
        text = str(value)
        return text if len(text) <= limit else text[: limit - 1] + "…"

    def empty(text: str = "لا توجد بيانات في هذه الفترة") -> Paragraph:
        return Paragraph(_ar(text), note)

    story = [
        Paragraph(_ar("تقرير إسناد — قاعدة البيانات والموقع"), title),
        Paragraph(_ar(f"الفترة: من {summary.date_from} إلى {summary.date_to} · "
                      f"أُنشئ في {datetime.now(UTC):%Y-%m-%d %H:%M} UTC"), note),
        Spacer(1, 0.4 * cm),
        Paragraph(_ar("الملخص"), heading),
        table(["البند", "القيمة"], [list(r) for r in _summary_rows(summary)], [9 * cm, 7 * cm]),
    ]

    active_days = [d for d in summary.daily if d.visits or d.searches or d.chats]
    story.append(Paragraph(_ar("النشاط اليومي"), heading))
    story.append(
        table(["اليوم", "الزيارات", "الزوار", "البحث", "الحوار"],
              [[d.day, d.visits, d.visitors, d.searches, d.chats] for d in active_days],
              [4 * cm, 3 * cm, 3 * cm, 3 * cm, 3 * cm])
        if active_days else empty()
    )

    for label, items in (("أكثر عمليات البحث", summary.top_searches),
                         ("أكثر أسئلة الحوار", summary.top_questions),
                         ("نتائج البحث", summary.search_results)):
        story.append(Paragraph(_ar(label), heading))
        story.append(table(["النص", "العدد"], [[short(i.label), i.count] for i in items],
                           [13.5 * cm, 2.5 * cm]) if items else empty())

    story.append(Paragraph(_ar("الملفات في قاعدة البيانات"), heading))
    story.append(
        table(["اسم الملف", "النوع", "الحجم", "المقاطع", "تاريخ الرفع"],
              [[short(s.filename, 45), KIND_LABELS.get(s.kind, s.kind), _size(s.size_bytes), s.chunks,
                s.uploaded_at[:10]] for s in sources],
              [6.5 * cm, 3 * cm, 2 * cm, 2 * cm, 2.5 * cm])
        if sources else empty("لا توجد ملفات مرفوعة")
    )

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, rightMargin=1.8 * cm, leftMargin=1.8 * cm,
        topMargin=1.6 * cm, bottomMargin=1.6 * cm, title="Isnad report",
    ).build(story)
    return buf.getvalue()
