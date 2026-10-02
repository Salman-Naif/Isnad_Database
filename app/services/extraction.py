"""
Text extraction from source files: PDF, Word, plain text and scanned documents.

  - PDF: every page is opened with PDFium (the engine inside Chrome), which reads any PDF,
    including scans compressed with JBIG2/JPX/CCITT that pure-Python readers can't decode.
    A page's text layer is used when it is real text; a page without one (a scan) or with
    an unusable one (garbled automatic OCR, broken font encoding, reversed Arabic) is
    rendered at OCR_DPI and read with OCR.
  - Word: paragraphs and tables, plus OCR of scanned pages pasted into the document.
  - Plain text: UTF-8, UTF-16 or Windows Arabic (cp1256).
  - Images (.png, .jpg, .tif, ...): OCR, every frame of a multi-page TIFF.

OCR itself (vision model with a second reading and a referee, or Tesseract) lives in
app.services.ocr. Pages are read in parallel (OCR_CONCURRENCY) and kept in order.

Works on raw bytes so the same code serves both web uploads and the CLI scripts.
"""

import io
import logging
import re
import unicodedata
from collections.abc import Callable, Iterable, Iterator
from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
from dataclasses import dataclass
from pathlib import Path

from app.config import get_settings
from app.services.ocr import OcrStats, OcrUnavailable, PageReader, PageResult, vision_enabled

# Receives a short Arabic status line, e.g. "قراءة الصفحات 45/300".
Progress = Callable[[str], None]

logger = logging.getLogger(__name__)

PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
TEXT_EXTENSIONS = {".txt", ".md"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
SUPPORTED_EXTENSIONS = PDF_EXTENSIONS | DOCX_EXTENSIONS | TEXT_EXTENSIONS | IMAGE_EXTENSIONS

# A PDF page with fewer extracted characters than this is treated as scanned.
MIN_TEXT_LAYER_CHARS = 20
# Rendered pages are capped at this many pixels on the long side (oversized scans).
RENDER_MAX_LONG_SIDE = 3600
# Images pasted into a Word file smaller than this (logos, icons) are not OCR'd.
DOCX_MIN_SCAN_SIDE = 800

_ARABIC = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿ]")
_LATIN = re.compile(r"[A-Za-z]")
_ARABIC_WORD = re.compile(r"[ء-ي]+")


class ExtractionError(Exception):
    """The file can't be read: unsupported type, corrupt file, or OCR unavailable."""


@dataclass
class ExtractedText:
    text: str
    ocr_pages: int = 0  # pages/images that were read with OCR
    note: str | None = None  # pages worth a human look, shown next to the file


def latin_share(text: str) -> float:
    """Share of Latin letters among Arabic + Latin letters (0 = all Arabic, 1 = all Latin)."""
    arabic = len(_ARABIC.findall(text))
    latin = len(_LATIN.findall(text))
    return latin / (arabic + latin) if arabic + latin else 0.0


def looks_garbled(text: str) -> bool:
    """Arabic text interleaved with stray Latin letters — the signature of bad OCR.

    Arabic hadith sources contain almost no Latin script, so a real mix of the two means
    Arabic glyphs were misread (e.g. "BE gall", "INKS"). Pure Latin text (an English page)
    and pure Arabic text both pass.
    """
    if len(_ARABIC.findall(text)) + len(_LATIN.findall(text)) < 20:
        return False
    return 0.08 < latin_share(text) < 0.9


def looks_broken(text: str) -> bool:
    """A PDF text layer whose Arabic is unusable although its letters are Arabic.

    Two common faults of PDFs made by old software:
      - letters split apart ("ا ل ح د ي ث") — the font's glyphs carry no word spacing info
      - words stored backwards ("ثيدحلا") — visual order instead of logical order; the
        article "ال" then shows up at the end of words ("...لا") instead of the start
    """
    words = _ARABIC_WORD.findall(text)
    if len(words) < 20:
        return False
    if sum(len(w) == 1 for w in words) / len(words) > 0.35:
        return True
    starts = sum(w.startswith("ال") and len(w) > 3 for w in words)
    ends = sum(w.endswith("لا") and len(w) > 3 for w in words)
    return ends >= 5 and ends > starts * 2


def text_layer_usable(text: str) -> bool:
    return (
        len(text.strip()) >= MIN_TEXT_LAYER_CHARS
        and not looks_garbled(text)
        and not looks_broken(text)
    )


def normalize_forms(text: str) -> str:
    """Turn Arabic presentation forms (ﻻ ﺍ ﺑ …) that PDFs often contain into normal letters.

    Without this, a word extracted from a PDF and the same word typed by a visitor are
    different characters and never match.
    """
    return unicodedata.normalize("NFKC", text)


def extract_text(data: bytes, filename: str, progress: Progress | None = None) -> ExtractedText:
    """Extract plain text from a file's bytes, picking the reader by extension."""
    ext = Path(filename).suffix.lower()
    progress = progress or (lambda _: None)
    if ext in PDF_EXTENSIONS:
        return _from_pdf(data, progress)
    if ext in DOCX_EXTENSIONS:
        return _from_docx(data, progress)
    if ext in TEXT_EXTENSIONS:
        return ExtractedText(_from_txt(data))
    if ext in IMAGE_EXTENSIONS:
        return _from_image(data, progress)
    raise ExtractionError(f"نوع الملف غير مدعوم: {ext or filename}")


# --- Reading a sequence of text blocks and scanned images, OCR in parallel ---


@dataclass
class Scan:
    """An image to OCR. `layer` is the page's own (unusable) text, kept if OCR finds nothing."""

    image: object
    number: int  # page number, for the review note
    layer: str = ""


def _read_blocks(
    blocks: Iterable[str | Scan], progress: Progress, total: int | None, label: str
) -> ExtractedText:
    """Text blocks are kept as they are; Scans are OCR'd in parallel. Order is preserved."""
    settings = get_settings()
    texts: dict[int, str] = {}
    stats = OcrStats()
    window = settings.ocr_concurrency * 2  # rendered pages waiting in memory at most

    def report() -> None:
        count = f" {len(texts)}/{total}" if total else ""
        ocr_note = f" — منها {stats.pages} بالتعرّف الضوئي" if stats.pages else ""
        progress(f"{label}{count}{ocr_note}")

    def settle(pending: dict[Future, tuple[int, Scan]], until) -> None:
        finished, _ = wait(list(pending), return_when=until)
        for future in finished:
            index, scan = pending.pop(future)
            result: PageResult = future.result()
            stats.add(scan.number, result)
            texts[index] = _choose(result.text, scan.layer)
        report()

    pool = ThreadPoolExecutor(settings.ocr_concurrency, thread_name_prefix="ocr")
    try:
        with PageReader() as reader:
            pending: dict[Future, tuple[int, Scan]] = {}
            for index, block in enumerate(blocks):
                if isinstance(block, Scan):
                    pending[pool.submit(reader.read, block.image)] = (index, block)
                    if len(pending) >= window:
                        settle(pending, FIRST_COMPLETED)
                else:
                    texts[index] = block
                    if not pending:
                        report()
            while pending:
                settle(pending, FIRST_COMPLETED)
    except OcrUnavailable as exc:
        raise ExtractionError(str(exc)) from exc
    finally:
        pool.shutdown(wait=False, cancel_futures=True)

    if stats.pages:
        logger.info(
            "OCR: %d pages, $%.4f, %d disputed, %d via Tesseract fallback",
            stats.pages, stats.cost, len(stats.disputed_pages), len(stats.fallback_pages),
        )
    text = "\n\n".join(t for _, t in sorted(texts.items()) if t.strip())
    return ExtractedText(text, ocr_pages=stats.pages, note=stats.note())


def _choose(ocr_text: str, layer: str) -> str:
    """OCR replaces an unusable text layer — unless OCR found nothing, or (Tesseract only)
    its reading is even more garbled than the layer."""
    if not ocr_text.strip():
        return layer
    if layer.strip() and not vision_enabled() and latin_share(ocr_text) > latin_share(layer):
        return layer
    return ocr_text


# --- PDF ---


def _from_pdf(data: bytes, progress: Progress) -> ExtractedText:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        raise ExtractionError(f"تعذّرت قراءة ملف PDF (تالف أو محمي بكلمة مرور): {exc}") from exc

    try:
        total = len(pdf)
        return _read_blocks(_pdf_pages(pdf), progress, total, "قراءة الصفحات")
    finally:
        pdf.close()


def _pdf_pages(pdf) -> Iterator[str | Scan]:
    """Each page's text layer if usable, otherwise its rendered image."""
    dpi = get_settings().ocr_dpi
    for index in range(len(pdf)):
        page = pdf[index]
        try:
            try:
                textpage = page.get_textpage()
                layer = normalize_forms(textpage.get_text_bounded())
                textpage.close()
            except Exception:
                logger.warning("Text layer unreadable on PDF page %d", index + 1)
                layer = ""
            if text_layer_usable(layer):
                yield layer
                continue
            width, height = page.get_size()  # points (1/72 inch)
            scale = min(dpi / 72, RENDER_MAX_LONG_SIDE / max(width, height, 1))
            # copy(): the rendered bitmap's memory belongs to PDFium
            image = page.render(scale=scale, grayscale=True).to_pil().copy()
        finally:
            page.close()
        yield Scan(image, index + 1, layer)


# --- Word ---


def _from_docx(data: bytes, progress: Progress) -> ExtractedText:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"تعذّرت قراءة ملف Word: {exc}") from exc
    return _read_blocks(_docx_blocks(doc), progress, None, "قراءة المستند")


def _docx_blocks(doc) -> Iterator[str | Scan]:
    """Paragraphs in order, each followed by the scans pasted in it, then the tables."""
    from docx.oxml.ns import qn
    from PIL import Image

    scans = 0
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            yield paragraph.text
        for blip in paragraph._p.iter(qn("a:blip")):
            part = doc.part.related_parts.get(blip.get(qn("r:embed")))
            if part is None or not part.content_type.startswith("image/"):
                continue
            try:
                image = Image.open(io.BytesIO(part.blob))
                image.load()
            except Exception:
                logger.warning("Unreadable image in Word file")
                continue
            if max(image.size) >= DOCX_MIN_SCAN_SIDE:
                scans += 1
                yield Scan(image, scans)
    for table in doc.tables:
        for row in table.rows:
            line = " | ".join(cell.text.strip() for cell in row.cells)
            if line.replace("|", "").strip():
                yield line


# --- Plain text ---


def _from_txt(data: bytes) -> str:
    # utf-8-sig also reads plain UTF-8; cp1256 is the legacy Windows Arabic encoding.
    for encoding in ("utf-8-sig", "utf-16", "cp1256"):
        if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# --- Scanned images ---


def _from_image(data: bytes, progress: Progress) -> ExtractedText:
    from PIL import Image, ImageSequence

    try:
        image = Image.open(io.BytesIO(data))
        frames = [frame.copy() for frame in ImageSequence.Iterator(image)]  # multi-page TIFF
    except Exception as exc:
        raise ExtractionError(f"تعذّرت قراءة الصورة: {exc}") from exc
    scans = [Scan(frame, number) for number, frame in enumerate(frames, start=1)]
    return _read_blocks(scans, progress, len(scans), "قراءة الصور")
